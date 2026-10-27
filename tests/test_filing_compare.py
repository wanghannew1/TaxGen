"""报税后核对（生成文件 vs 个税端导出）单元测试。"""
import csv
import os

import pytest

import filing_compare
import filing_history
import tax_return


@pytest.fixture()
def db(tmp_path, monkeypatch):
    path = str(tmp_path / "filing_compare.db")
    monkeypatch.setattr(tax_return, "DB_PATH", path)
    monkeypatch.setattr(filing_history, "DB_PATH", path)
    monkeypatch.setattr(filing_compare, "DB_PATH", path)
    filing_compare.init_db()
    filing_history.init_db()
    return path


def _write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    return str(path)


def _gen_file(tmp_path):
    header = ["工号", "*姓名", "*证件类型", "*证件号码", "本期收入", "本期免税收入",
              "基本养老保险费", "基本医疗保险费", "失业保险费", "住房公积金", "备注"]
    rows = [
        ["G1", "张三", "居民身份证", "110101199001010011", 5000, 0, 500, 100, 50, 200, "单位A"],
        ["G2", "李四", "居民身份证", "110101199001010022", 6000, 0, 600, 120, 60, 240, "单位A"],
        ["G3", "王五", "居民身份证", "110101199001010033", 7000, 0, 700, 140, 70, 280, "单位A"],
        ["G3", "王五", "居民身份证", "110101199001010033", 3000, 0, 300, 60, 30, 120, "单位B"],
        ["G4", "赵六", "居民身份证", "110101199001010044", 8000, 0, 800, 160, 80, 320, "单位A"],
        ["G5", "孙七", "居民身份证", "110101199001010077", 4500, 0, 450, 90, 45, 180, "单位B"],
    ]
    return _write_csv(tmp_path / "gen_normal.csv", header, rows)


def _filing_file(tmp_path):
    # 个税端税款计算导出（含 所得期间起/税款所属期起），用于验证「生成文件导入」会拒绝
    header = ["工号", "姓名", "证件类型", "证件号码", "税款所属期起", "税款所属期止",
              "所得项目", "本期收入", "本期免税收入", "本期基本养老保险费",
              "本期基本医疗保险费", "本期失业保险费", "本期住房公积金",
              "累计应扣缴税额", "已缴税额", "应补(退)税额", "备注"]
    rows = [
        ["A1", "张三", "居民身份证", "110101199001010011", "2026-09-01", "2026-09-30",
         "正常工资薪金", 5000, 0, 500, 100, 50, 200, 10, 0, 10, "单位A"],
    ]
    return _write_csv(tmp_path / "filing_tax_csv.csv", header, rows)


def _filing_records():
    def rec(cert, name, income, remark="单位A", tax_due=10):
        return {
            "cert_no": cert, "name": name, "emp_no": "", "id_type": "居民身份证",
            "month": 202609, "item_type": "税款计算", "income": income, "tax_free": 0,
            "pension": 500, "medical": 100, "unemployment": 50, "housing": 200,
            "insurance": 850, "tax_accum": 0, "tax_paid": 0, "tax_due": tax_due,
            "remark": remark, "raw": "{}", "source_file": "test",
        }

    return [
        rec("110101199001010011", "张三", 5000, "单位A", 10),
        rec("110101199001010022", "李四", 6200, "单位A", 20),   # 收入比生成多 200
        rec("110101199001010055", "钱七", 4000, "单位A", 5),
        rec("110101199001010088", "孙七", 4500, "单位B", 10),
    ]


def test_parse_generated_normal(db, tmp_path):
    path = _gen_file(tmp_path)
    records = filing_compare.parse_generated_file(path, 202609)
    assert len(records) == 6
    assert records[0]["item_type"] == "正常工资薪金"
    assert records[0]["insurance"] == 850
    assert records[0]["month"] == 202609
    # 重复证件保留两行
    assert sum(1 for r in records if r["cert_no"] == "110101199001010033") == 2


def test_parse_generated_requires_month(db, tmp_path):
    path = _gen_file(tmp_path)
    with pytest.raises(ValueError, match="月份"):
        filing_compare.parse_generated_file(path, 0)


def test_parse_generated_rejects_filing_export(db, tmp_path):
    path = _filing_file(tmp_path)
    with pytest.raises(ValueError, match="个税端导出"):
        filing_compare.parse_generated_file(path, 202609)


def test_compare_statuses_and_same_name(db, tmp_path):
    gen_path = _gen_file(tmp_path)
    records = filing_compare.parse_generated_file(gen_path, 202609)
    filing_compare.import_generated_records(records)
    filing_history.import_filing_records(_filing_records())

    result = filing_compare.compare_generated_vs_filing(
        202609, "正常工资薪金", "税款计算")
    by_cert = {d["cert_no"]: d for d in result["details"]}

    assert by_cert["110101199001010011"]["status"] == "一致"
    assert by_cert["110101199001010022"]["status"] == "收入差异"
    assert by_cert["110101199001010033"]["status"] == "生成侧重复"
    assert by_cert["110101199001010044"]["status"] == "仅生成"
    assert by_cert["110101199001010055"]["status"] == "仅个税端"
    assert by_cert["110101199001010077"]["status"] == "仅生成"
    assert by_cert["110101199001010088"]["status"] == "仅个税端"
    assert by_cert["110101199001010077"]["same_name_flag"] is True
    assert by_cert["110101199001010088"]["same_name_flag"] is True

    s = result["summary"]
    assert s["gen_certs"] == 5
    assert s["gen_rows"] == 6
    assert s["fil_count"] == 4
    assert s["matched"] == 1
    assert s["diff"] == 2
    assert s["gen_only"] == 2
    assert s["fil_only"] == 2
    assert s["dup_certs"] == 1
    assert s["same_name_groups"] == 1

    same = result["same_name"][0]
    assert same["name"] == "孙七"
    assert "疑似挂错" in same["type"]
    assert same["gen_certs"] == "110101199001010077"
    assert same["fil_certs"] == "110101199001010088"


def test_import_replaces_same_month_type(db, tmp_path):
    path = _gen_file(tmp_path)
    filing_compare.import_generated_records(
        filing_compare.parse_generated_file(path, 202609))
    types = filing_compare.get_generated_types(202609)
    assert types == [{"item_type": "正常工资薪金", "count": 6}]

    # 再导入同月同类型：覆盖，不累加
    filing_compare.import_generated_records(
        filing_compare.parse_generated_file(path, 202609))
    types = filing_compare.get_generated_types(202609)
    assert types == [{"item_type": "正常工资薪金", "count": 6}]


def test_export_comparison(db, tmp_path):
    from openpyxl import load_workbook

    gen_path = _gen_file(tmp_path)
    filing_compare.import_generated_records(
        filing_compare.parse_generated_file(gen_path, 202609))
    filing_history.import_filing_records(_filing_records())
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    filename = filing_compare.export_comparison(
        202609, "正常工资薪金", "税款计算", str(out_dir))
    path = out_dir / filename
    assert path.exists()
    wb = load_workbook(path)
    assert set(wb.sheetnames) >= {"汇总", "比对明细", "同名不同证", "生成侧重复"}
    ws = wb["比对明细"]
    assert ws.max_row == 1 + 7  # 表头 + 7 个证件（生成 5 + 个税端独有 2）
    assert wb["同名不同证"].max_row >= 2


def test_api_import_generated_and_compare(db, tmp_path, app_client):
    path = _gen_file(tmp_path)
    with open(path, "rb") as f:
        resp = app_client.post(
            "/api/filing/import-generated",
            data={"month": "202609", "files": (f, "gen_normal.csv")},
            content_type="multipart/form-data",
        )
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["count"] == 6
    assert body["types"] == ["正常工资薪金"]

    filing_history.import_filing_records(_filing_records())
    resp = app_client.get(
        "/api/filing/compare?month=202609&gen_item_type=正常工资薪金&filing_item_type=税款计算")
    assert resp.status_code == 200
    body = resp.get_json()
    assert body["summary"]["gen_certs"] == 5
    assert body["summary"]["dup_certs"] == 1
    assert body["total"] == 7
    assert body["same_name"][0]["name"] == "孙七"
